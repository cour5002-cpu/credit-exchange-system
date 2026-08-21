import csv
import io
import re
from datetime import date
from dataclasses import dataclass
from uuid import uuid4
from zipfile import BadZipFile

from openpyxl.utils.exceptions import InvalidFileException
from openpyxl import load_workbook

from app.extensions import db
from app.models.operation_log import OperationLog
from app.models.student import Student
from app.models.student_hour_account import StudentHourAccount
from app.models.teacher import Teacher
from app.models.user import User


DEFAULT_PASSWORD = "123456"
EMAIL_PATTERN = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
USER_STATUSES = {"active", "disabled"}
STUDENT_STATUSES = {"active", "disabled", "graduated"}
TEACHER_STATUSES = {"active", "disabled"}
ROLE_FLAG_MAPPING = {
    "指导老师": "advisor",
    "审核老师": "reviewer",
    "advisor": "advisor",
    "reviewer": "reviewer",
}


@dataclass
class ImportErrorItem:
    row_no: int
    field: str | None
    message: str

    def to_dict(self):
        return {
            "row_no": self.row_no,
            "field": self.field,
            "message": self.message,
        }


def read_import_rows(file_storage):
    filename = (file_storage.filename or "").lower()
    try:
        if filename.endswith(".csv"):
            return read_csv_rows(file_storage)
        if filename.endswith((".xlsx", ".xlsm")):
            return read_excel_rows(file_storage)
        if filename.endswith(".xls"):
            return read_xls_rows(file_storage)
    except (UnicodeDecodeError, InvalidFileException, BadZipFile) as exc:
        raise ValueError("导入文件无法解析，请检查文件格式") from exc
    raise ValueError("仅支持 .xlsx、.xlsm、.xls 或 .csv 文件")


def read_csv_rows(file_storage):
    content = file_storage.read().decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(content))
    rows = []
    for index, row in enumerate(reader, start=2):
        normalized = {
            (key or "").strip(): (value or "").strip()
            for key, value in row.items()
        }
        rows.append((index, normalized))
    return rows


def read_excel_rows(file_storage):
    workbook = load_workbook(file_storage, read_only=True, data_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    headers = next(rows_iter, None)
    if not headers:
        return []

    normalized_headers = [_cell_to_text(header) for header in headers]
    rows = []
    for index, values in enumerate(rows_iter, start=2):
        if not values or all(_cell_to_text(value) == "" for value in values):
            continue
        row = {
            header: _cell_to_text(value)
            for header, value in zip(normalized_headers, values)
            if header
        }
        rows.append((index, row))
    return rows


def read_xls_rows(file_storage):
    try:
        import xlrd
    except ImportError as exc:
        raise ValueError("当前环境缺少 xlrd，无法解析 .xls 文件") from exc

    try:
        workbook = xlrd.open_workbook(file_contents=file_storage.read())
    except Exception as exc:
        raise ValueError("导入文件无法解析，请检查文件格式") from exc
    sheet = workbook.sheet_by_index(0)
    if sheet.nrows == 0:
        return []

    headers = [_cell_to_text(sheet.cell_value(0, col_index)) for col_index in range(sheet.ncols)]
    rows = []
    for row_index in range(1, sheet.nrows):
        values = [sheet.cell_value(row_index, col_index) for col_index in range(sheet.ncols)]
        if not values or all(_cell_to_text(value) == "" for value in values):
            continue
        row = {
            header: _cell_to_text(value)
            for header, value in zip(headers, values)
            if header
        }
        rows.append((row_index + 1, row))
    return rows


def import_students(file_storage, mode="upsert", imported_by_user_id=None, ip=None):
    batch_no = f"students-{uuid4().hex[:12]}"
    success_count = 0
    errors: list[ImportErrorItem] = []
    seen_student_numbers = set()
    seen_usernames = set()

    rows = read_import_rows(file_storage)
    errors.extend(_validate_required_headers(rows, ["student_no", "name"]))
    if errors:
        _finish_import(errors, imported_by_user_id, ip, "students", batch_no, success_count)
        return _result(success_count, errors, batch_no)

    for row_no, row in rows:
        student_no = row.get("student_no") or row.get("学号")
        name = row.get("name") or row.get("姓名")
        if not student_no:
            errors.append(ImportErrorItem(row_no, "student_no", "学号不能为空"))
            continue
        if not name:
            errors.append(ImportErrorItem(row_no, "name", "姓名不能为空"))
            continue
        if student_no in seen_student_numbers:
            errors.append(ImportErrorItem(row_no, "student_no", "学号在本次导入文件中重复"))
            continue
        seen_student_numbers.add(student_no)
        email = row.get("email") or row.get("邮箱")
        if email and not _is_valid_email(email):
            errors.append(ImportErrorItem(row_no, "email", "邮箱格式不正确"))
            continue
        student_status = row.get("status") or row.get("状态") or "active"
        if student_status not in STUDENT_STATUSES:
            errors.append(ImportErrorItem(row_no, "status", "学生状态只能是 active、disabled 或 graduated"))
            continue
        graduation_text = row.get("expected_graduation_date") or row.get("预计毕业日期")
        if graduation_text:
            try:
                expected_graduation_date = date.fromisoformat(graduation_text[:10])
            except ValueError:
                errors.append(ImportErrorItem(row_no, "expected_graduation_date", "预计毕业日期格式必须为 YYYY-MM-DD"))
                continue
        else:
            expected_graduation_date = None

        existing_student = Student.query.filter_by(student_no=student_no).first()
        username = row.get("username") or row.get("账号") or student_no
        if username in seen_usernames:
            errors.append(ImportErrorItem(row_no, "username", "账号在本次导入文件中重复"))
            continue
        seen_usernames.add(username)
        existing_user = User.query.filter_by(username=username).first()
        if mode == "append" and existing_student:
            errors.append(ImportErrorItem(row_no, "student_no", "学号已存在"))
            continue
        if mode == "append" and existing_user:
            errors.append(ImportErrorItem(row_no, "username", "登录账号已存在"))
            continue
        if existing_user and (not existing_student or existing_student.user_id != existing_user.id):
            errors.append(ImportErrorItem(row_no, "username", "登录账号已存在"))
            continue

        user = existing_user or User(username=username, role="student", real_name=name, status="active")
        user.real_name = name
        user.role = "student"
        user.phone = row.get("phone") or row.get("手机号") or user.phone
        user.email = email or user.email
        user.status = "disabled" if student_status in {"disabled", "graduated"} else "active"
        if not existing_user:
            user.set_password(row.get("password") or row.get("密码") or DEFAULT_PASSWORD)
            db.session.add(user)
            db.session.flush()

        student = existing_student or Student(user_id=user.id, student_no=student_no, name=name)
        student.user_id = user.id
        student.name = name
        student.college = row.get("college") or row.get("学院") or "管理学院"
        student.major = row.get("major") or row.get("专业") or ""
        student.grade = row.get("grade") or row.get("年级") or ""
        student.class_name = row.get("class_name") or row.get("班级") or ""
        student.expected_graduation_date = expected_graduation_date
        student.status = student_status
        student.import_batch_no = batch_no
        if not existing_student:
            db.session.add(student)
            db.session.flush()
        if not StudentHourAccount.query.filter_by(student_id=student.id).first():
            db.session.add(
                StudentHourAccount(
                    student_id=student.id,
                    total_earned_hours=0,
                    total_exchanged_hours=0,
                    available_hours=0,
                )
            )
        success_count += 1

    _finish_import(errors, imported_by_user_id, ip, "students", batch_no, success_count)
    return _result(success_count, errors, batch_no)


def import_teachers(file_storage, mode="upsert", imported_by_user_id=None, ip=None):
    batch_no = f"teachers-{uuid4().hex[:12]}"
    success_count = 0
    errors: list[ImportErrorItem] = []
    seen_teacher_numbers = set()
    seen_usernames = set()

    rows = read_import_rows(file_storage)
    errors.extend(_validate_required_headers(rows, ["teacher_no", "name"]))
    if errors:
        _finish_import(errors, imported_by_user_id, ip, "teachers", batch_no, success_count)
        return _result(success_count, errors, batch_no)

    for row_no, row in rows:
        teacher_no = row.get("teacher_no") or row.get("工号")
        name = row.get("name") or row.get("姓名")
        if not teacher_no:
            errors.append(ImportErrorItem(row_no, "teacher_no", "工号不能为空"))
            continue
        if not name:
            errors.append(ImportErrorItem(row_no, "name", "姓名不能为空"))
            continue
        if teacher_no in seen_teacher_numbers:
            errors.append(ImportErrorItem(row_no, "teacher_no", "工号在本次导入文件中重复"))
            continue
        seen_teacher_numbers.add(teacher_no)
        email = row.get("email") or row.get("邮箱")
        if email and not _is_valid_email(email):
            errors.append(ImportErrorItem(row_no, "email", "邮箱格式不正确"))
            continue
        teacher_status = row.get("status") or row.get("状态") or "active"
        if teacher_status not in TEACHER_STATUSES:
            errors.append(ImportErrorItem(row_no, "status", "教师状态只能是 active 或 disabled"))
            continue
        try:
            role_flags = _normalize_role_flags(row.get("role_flags") or row.get("角色能力"))
        except ValueError as exc:
            errors.append(ImportErrorItem(row_no, "role_flags", str(exc)))
            continue

        existing_teacher = Teacher.query.filter_by(teacher_no=teacher_no).first()
        username = row.get("username") or row.get("账号") or teacher_no
        if username in seen_usernames:
            errors.append(ImportErrorItem(row_no, "username", "账号在本次导入文件中重复"))
            continue
        seen_usernames.add(username)
        existing_user = User.query.filter_by(username=username).first()
        if mode == "append" and existing_teacher:
            errors.append(ImportErrorItem(row_no, "teacher_no", "工号已存在"))
            continue
        if mode == "append" and existing_user:
            errors.append(ImportErrorItem(row_no, "username", "登录账号已存在"))
            continue
        if existing_user and (not existing_teacher or existing_teacher.user_id != existing_user.id):
            errors.append(ImportErrorItem(row_no, "username", "登录账号已存在"))
            continue

        primary_role = role_flags[0] if len(role_flags) == 1 else "teacher"
        user = existing_user or User(username=username, role=primary_role, real_name=name, status="active")
        user.real_name = name
        user.role = primary_role
        user.phone = row.get("phone") or row.get("手机号") or user.phone
        user.email = email or user.email
        user.status = teacher_status
        if not existing_user:
            user.set_password(row.get("password") or row.get("密码") or DEFAULT_PASSWORD)
            db.session.add(user)
            db.session.flush()

        teacher = existing_teacher or Teacher(user_id=user.id, teacher_no=teacher_no, name=name)
        teacher.user_id = user.id
        teacher.name = name
        teacher.major_name = row.get("major_name") or row.get("专业方向") or teacher.major_name
        teacher.course_name = row.get("course_name") or row.get("课程名称") or teacher.course_name
        teacher.role_flags = ",".join(role_flags)
        teacher.status = teacher_status
        if not existing_teacher:
            db.session.add(teacher)
        success_count += 1

    _finish_import(errors, imported_by_user_id, ip, "teachers", batch_no, success_count)
    return _result(success_count, errors, batch_no)


def import_admins(file_storage, mode="upsert", imported_by_user_id=None, ip=None):
    batch_no = f"admins-{uuid4().hex[:12]}"
    success_count = 0
    errors: list[ImportErrorItem] = []
    seen_usernames = set()

    rows = read_import_rows(file_storage)
    errors.extend(_validate_required_headers(rows, ["username", "name"]))
    if errors:
        _finish_import(errors, imported_by_user_id, ip, "admins", batch_no, success_count)
        return _result(success_count, errors, batch_no)

    for row_no, row in rows:
        username = row.get("username") or row.get("账号")
        name = row.get("name") or row.get("姓名")
        if not username:
            errors.append(ImportErrorItem(row_no, "username", "账号不能为空"))
            continue
        if not name:
            errors.append(ImportErrorItem(row_no, "name", "姓名不能为空"))
            continue
        if username in seen_usernames:
            errors.append(ImportErrorItem(row_no, "username", "账号在本次导入文件中重复"))
            continue
        seen_usernames.add(username)
        email = row.get("email") or row.get("邮箱")
        if email and not _is_valid_email(email):
            errors.append(ImportErrorItem(row_no, "email", "邮箱格式不正确"))
            continue
        user_status = row.get("status") or row.get("状态") or "active"
        if user_status not in USER_STATUSES:
            errors.append(ImportErrorItem(row_no, "status", "管理员账号状态只能是 active 或 disabled"))
            continue

        user = User.query.filter_by(username=username).first()
        if mode == "append" and user:
            errors.append(ImportErrorItem(row_no, "username", "账号已存在"))
            continue
        if user and user.role != "admin":
            errors.append(ImportErrorItem(row_no, "username", "账号已被非管理员使用"))
            continue
        if not user:
            user = User(username=username, role="admin", real_name=name, status="active")
            user.set_password(row.get("password") or row.get("密码") or DEFAULT_PASSWORD)
            db.session.add(user)
        user.real_name = name
        user.role = "admin"
        user.phone = row.get("phone") or row.get("手机号") or user.phone
        user.email = email or user.email
        user.status = user_status
        success_count += 1

    _finish_import(errors, imported_by_user_id, ip, "admins", batch_no, success_count)
    return _result(success_count, errors, batch_no)


def _normalize_role_flags(raw_value):
    raw_items = [item.strip() for item in (raw_value or "advisor,reviewer").replace("；", ",").split(",")]
    flags = []
    for item in raw_items:
        if not item:
            continue
        normalized = ROLE_FLAG_MAPPING.get(item)
        if not normalized:
            raise ValueError("教师角色能力只能是 advisor、reviewer、指导老师或审核老师")
        if normalized and normalized not in flags:
            flags.append(normalized)
    return flags or ["advisor", "reviewer"]


def _is_valid_email(value):
    return bool(EMAIL_PATTERN.match(value.strip()))


def _validate_required_headers(rows, required_fields):
    if not rows:
        return [ImportErrorItem(1, None, "导入文件没有可读取的数据行")]

    headers = set(rows[0][1].keys())
    errors = []
    for field in required_fields:
        if field not in headers and _chinese_header(field) not in headers:
            errors.append(ImportErrorItem(1, field, f"缺少必需列：{field}"))
    return errors


def _chinese_header(field):
    mapping = {
        "student_no": "学号",
        "teacher_no": "工号",
        "username": "账号",
        "name": "姓名",
    }
    return mapping.get(field)


def _finish_import(errors, imported_by_user_id, ip, target, batch_no, success_count):
    if errors:
        db.session.rollback()
        return
    if imported_by_user_id:
        db.session.add(
            OperationLog(
                user_id=imported_by_user_id,
                module="base_data",
                biz_type=f"import_{target}",
                action="import",
                detail=f"batch_no={batch_no}; success_count={success_count}",
                ip=ip,
            )
        )
    db.session.commit()


def _cell_to_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _result(success_count, errors, batch_no):
    return {
        "success_count": 0 if errors else success_count,
        "failed_count": len(errors),
        "batch_no": None if errors else batch_no,
        "errors": [error.to_dict() for error in errors],
    }
